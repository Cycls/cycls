"""The model's specs, decks, documents, slides and edit ops made ready for the design service.
Every `_prepare_*` returns what to send, or an error the model can act on."""
import json, pathlib, re
from ..paths import _resolve_path
from .brand import _HEX, _contrast, _first_color, _load_brand, _norm_hex, _readable_on
from .images import _BRAND_LOGOS, _DESIGN_IMAGES_MAX, _DESIGN_SVG_MAX, _image_slot, _place_image


# preset `size` names; a4-poster is 150dpi, so the default @2x render prints at 300dpi
_DESIGN_SIZES = {
    "square": [1080, 1080], "post-portrait": [1080, 1350],
    "story": [1080, 1920], "reel": [1080, 1920],
    "slide": [1920, 1080], "wide": [1920, 1080],
    "x-post": [1600, 900], "a4-poster": [1240, 1754],
}


_NUMERIC = ("x", "y", "w", "h", "size", "radius", "lineHeight", "letterSpacing",
            "opacity", "strokeWeight", "rotation", "gap", "width", "blur", "backdropBlur")


def _num(v):
    """A number the model may have quoted ("1500", "96px") → int/float, else None."""
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return v
    if isinstance(v, str) and (m := re.fullmatch(r"\s*(-?\d+(?:\.\d+)?)\s*(?:px)?\s*", v)):
        f = float(m.group(1))
        return int(f) if f.is_integer() else f
    return None


def _frame_size(fr):
    """A frame's [W, H] from any shape the model gave it (default 1080²); ValueError otherwise."""
    s = fr.get("size")
    if s is None and "width" in fr and "height" in fr:
        s = [fr.pop("width"), fr.pop("height")]
    if isinstance(s, dict):
        s = [s.get("w", s.get("width")), s.get("h", s.get("height"))]
    if s is None:
        return [1080, 1080]
    if isinstance(s, str):
        key = s.strip().lower()
        if key in _DESIGN_SIZES:
            return list(_DESIGN_SIZES[key])
        if m := re.fullmatch(r"(\d+)\s*[x×]\s*(\d+)", key):
            return [int(m.group(1)), int(m.group(2))]
        raise ValueError(f"unknown size {s!r} — use [W, H] or one of: {', '.join(_DESIGN_SIZES)}")
    if isinstance(s, (list, tuple)) and len(s) == 2 and all((_num(v) or 0) > 0 for v in s):
        return [_num(v) for v in s]
    raise ValueError(f"`size` {s!r} isn't a size — use [W, H] or a preset ({', '.join(_DESIGN_SIZES)})")


def _prepare_spec(spec, brand, root=None):
    """A spec made safe to render — sizes resolved, untyped nodes typed (the renderer drops them),
    images placed, brand defaults filled, never over a colour set — on a copy → (spec, error, notes)."""
    if isinstance(spec.get("deck"), dict):
        return _prepare_deck(spec, brand, root)
    if spec.get("pages") is not None:
        return _prepare_pages(spec, brand, root)
    colors = {_norm_hex(m) for m in _HEX.findall(json.dumps(spec))}
    spec = json.loads(json.dumps(spec))
    frames = spec["frames"] if isinstance(spec.get("frames"), list) and spec["frames"] else [spec]
    filled, branded_fonts, image_bytes = 0, 0, 0
    shapes = []                 # notes on nodes reinterpreted, once each
    first_size = None
    for i, fr in enumerate(frames, 1):
        if not isinstance(fr, dict):
            continue
        try:
            fr["size"] = W, H = _frame_size(fr)
        except ValueError as e:
            return None, f"Error: {e}.", []
        # PowerPoint takes slide 1's size for the whole file and letterboxes the rest
        if first_size is None:
            first_size = [W, H]
        elif [W, H] != first_size:
            return None, (f"Error: every slide of a deck must be the same size — slide 1 is "
                          f"{first_size[0]}×{first_size[1]}, slide {i} is {W}×{H}. Give every "
                          f"frame the same `size`."), []
        if brand and brand["primary"] and fr.get("fill") is None:
            fr["fill"], filled = brand["primary"], filled + 1
        bg = _first_color(fr.get("fill"))

        def prepare(n, top):
            """One node, made safe (a stack's children too). → an error, or None."""
            nonlocal filled, branded_fonts, image_bytes
            if n.get("type") is None and ("text" in n or "src" in n):
                n["type"] = "text" if "text" in n else "image"
            if n.get("type") in ("radial", "linear", "angular", "diamond") and isinstance(n.get("gradient"), list):
                # a gradient written as a node: a full-frame rect with that fill
                fill = {k: n[k] for k in ("gradient", "center", "radius", "angle") if n.get(k) is not None}
                fill["type"] = n["type"]
                keep = {k: n[k] for k in ("id", "opacity") if n.get(k) is not None}
                n.clear()
                n.update({"type": "rect", "x": 0, "y": 0, "w": W, "h": H, "fill": fill, **keep})
                shapes.append(f"A node of type \"{fill['type']}\" was read as a rect over the whole frame with that gradient as "
                              f"its `fill` — a gradient is a fill ({{\"gradient\": […], \"type\": \"{fill['type']}\"}}), not a node.")
            if n.get("type") == "list":
                items = n.get("items") if n.get("items") is not None else n.pop("text", None)
                if isinstance(items, str):
                    items = [line for line in items.splitlines() if line.strip()]
                if isinstance(items, list):
                    # {"text": …} items and typed bullet markers → plain strings
                    items = [i["text"] if isinstance(i, dict) and isinstance(i.get("text"), str) and "runs" not in i else i for i in items]
                    n["items"] = [re.sub(r"^\s*(?:[-–—•*·▪●]|\d{1,2}[.)])\s+", "", i).strip() if isinstance(i, str) else i for i in items]
            for k in _NUMERIC:                           # a string "1500" as y sinks the text
                if k in n and n[k] is not None and _num(n[k]) is None:
                    return f"Error: `{k}` must be a number, not {n[k]!r} ({json.dumps(n)[:80]})."
                if k in n and n[k] is not None:
                    n[k] = _num(n[k])
            if n.get("type") == "text":
                # the service resolves fonts; a brand kit fills unset ones (heading ≥48px, else body)
                if brand and n.get("font") is None:
                    face = brand["heading"] if (n.get("size") or 32) >= 48 else brand["body"]
                    face = face or brand["heading"] or brand["body"]
                    if face:
                        n["font"], branded_fonts = face, branded_fonts + 1
                if bg and n.get("color") is None and n.get("fill") is None:
                    n["color"] = _readable_on(bg)
            elif n.get("type") in ("rect", "ellipse", "line"):
                if n.get("fill") is None and n.get("color") is not None:
                    n["fill"] = n.pop("color")          # a shape draws only `fill`
                if brand and brand["accent"] and n.get("fill") is None and not n.get("stroke"):
                    n["fill"], filled = brand["accent"], filled + 1
            elif n.get("type") == "image":
                if n.pop("_frame", None) and top and W < _NO_FRAME:   # a library texture given no box
                    n["x"], n["y"] = n.get("x") or 0, n.get("y") or 0
                    n["w"], n["h"] = W - n["x"], H - n["y"]
                try:
                    image_bytes += _place_image(n, root)
                except ValueError as e:
                    return f"Error: {e}"
                if image_bytes > _DESIGN_IMAGES_MAX:
                    return (f"Error: the design's images total over {_DESIGN_IMAGES_MAX >> 20} MB — "
                            f"use smaller copies (longest side ~2000px).")
            elif n.get("type") == "pattern":
                # One SVG tile, which the service repeats over the node's box — the whole
                # frame when it gives none — in one ink, faint unless it says otherwise.
                if n.get("src") and not n.get("svg"):
                    path = _resolve_path(str(n.pop("src")), root)
                    if not path.is_file() or path.suffix.lower() != ".svg":
                        return f"Error: a pattern's `src` must be an .svg tile in the workspace ({json.dumps(n)[:80]})."
                    if path.stat().st_size > _DESIGN_SVG_MAX:
                        return f"Error: {path.name} is over {_DESIGN_SVG_MAX >> 10} KB — use a simpler tile."
                    n["svg"] = path.read_text(encoding="utf-8", errors="replace")
                if not (isinstance(n.get("svg"), str) and "<svg" in n["svg"]):
                    return ("Error: a pattern node needs `asset` (a pattern of the media library, e.g. \"asset\": \"moroccan\"), "
                            "`src` (an .svg tile in the workspace) or `svg` (the tile's markup).")
                if not (n.get("w") and n.get("h")):
                    if not top or W >= _NO_FRAME:
                        return "Error: this pattern needs `w` and `h` — the box it fills."
                    n["x"], n["y"] = n.get("x") or 0, n.get("y") or 0
                    n["w"], n["h"] = n.get("w") or W - n["x"], n.get("h") or H - n["y"]
                if bg and n.get("color") is None:
                    n["color"] = _readable_on(bg)
                if n.get("opacity") is None:
                    n["opacity"] = _PATTERN_OPACITY
            elif n.get("type") in ("icon", "svg", "qr", "list", "chart", "table"):
                kind = n["type"]
                if kind == "icon" and not (isinstance(n.get("name"), str) and ":" in n["name"]):
                    return f"Error: an icon needs `name` — an Iconify name like \"lucide:rocket\" or \"mdi:coffee\" ({json.dumps(n)[:80]})."
                if kind == "svg" and n.get("src") and not n.get("svg"):
                    path = _resolve_path(str(n.pop("src")), root)
                    if not path.is_file() or path.suffix.lower() != ".svg":
                        return f"Error: svg `src` must be an .svg file in the workspace ({json.dumps(n)[:80]})."
                    if path.stat().st_size > _DESIGN_SVG_MAX:
                        return f"Error: {path.name} is over {_DESIGN_SVG_MAX >> 10} KB — use a simpler SVG (or a PNG image)."
                    n["svg"] = path.read_text(encoding="utf-8", errors="replace")
                if kind == "svg" and not (isinstance(n.get("svg"), str) and "<svg" in n["svg"]):
                    return "Error: an svg node needs `svg` (the markup) or `src` (an .svg file in the workspace)."
                if kind == "qr" and not str(n.get("text") or "").strip():
                    return "Error: a qr node needs `text` — the link or text to encode."
                if kind == "list" and not isinstance(n.get("items"), list):
                    return "Error: a list needs `items` — a list of strings."
                if kind == "chart" and not isinstance((n.get("data") or {}).get("series") or n.get("values"), list):
                    return "Error: a chart needs `data`: {labels:[…], series:[{name, values:[…]}]}."
                if kind == "table" and not isinstance(n.get("rows"), list):
                    return "Error: a table needs `rows` (a list of rows, each a list of cells) and usually `columns`."
                if bg and n.get("color") is None and kind != "svg":
                    n["color"] = _readable_on(bg)
                if brand and n.get("font") is None and kind in ("list", "chart", "table") and brand["body"]:
                    n["font"], branded_fonts = brand["body"], branded_fonts + 1
            elif n.get("type") == "stack":
                kids = n.get("children")
                if not isinstance(kids, list) or not kids:
                    return f"Error: a stack needs `children` — the nodes it lays out ({json.dumps(n)[:80]})."
                for c in kids:
                    if not isinstance(c, dict):
                        return f"Error: a stack's children are nodes, not {c!r}."
                    if err := prepare(c, False):
                        return err
                # a child placed by hand: the stack lays it out after the one before — a label meant
                # for the shape before it (a button) lands beside it (horizontal) or under it
                placed = [str(c.get("id") or c.get("text") or c.get("type"))[:30] for c in kids
                          if isinstance(c, dict) and (c.get("x") is not None or c.get("y") is not None)]
                if placed and not any("are laid out one after another" in s for s in shapes):
                    shapes.append(f"In stack {json.dumps(str(n.get('id') or 'stack'))}, {', '.join(placed)} "
                                  f"{'has' if len(placed) == 1 else 'have'} its own x / y, which a stack doesn't read: its "
                                  f"children are laid out one after another ({'side by side' if n.get('direction') == 'horizontal' else 'one under the next'}). "
                                  f"To put a label ON a shape (a button, a badge), place both outside any stack in the same box — "
                                  f"the shape, then the text with the same x, y, w, h and \"align\": \"center\", \"valign\": \"center\".")
            else:
                return (f"Error: a node has type {n.get('type')!r} ({json.dumps(n)[:80]}) — "
                        f"every node needs a `type`: text, rect, ellipse, line, image, stack, list, "
                        f"table, chart, icon, svg, pattern or qr.")
            # outside the frame = laid out for another size; the renderer would pile it on the edge
            if top and n["type"] not in ("rect", "ellipse", "line") and ((n.get("x") or 0) >= W or (n.get("y") or 0) >= H):
                what = repr(n.get("text"))[:40] if n["type"] == "text" else f"a{'n' if n['type'] == 'image' else ''} {n['type']}"
                return (f"Error: {what} starts at ({n.get('x') or 0}, {n.get('y') or 0}), outside the "
                        f"{W}×{H} frame — set the frame's `size` (e.g. \"story\" for 1080×1920) "
                        f"or move it inside.")
            return None

        for n in fr.get("nodes") or []:
            if isinstance(n, dict) and (err := prepare(n, True)):
                return None, err, []
    notes = list(dict.fromkeys(shapes))
    if branded_fonts:
        notes.append(f"Brand fonts applied to {branded_fonts} text node(s) with no font "
                     f"(heading {brand['heading'] or brand['body']}, body {brand['body'] or brand['heading']}).")
    if brand and filled:
        notes.append(f"Brand kit applied to {filled} unset fill(s) "
                     f"(primary {brand['primary']}, accent {brand['accent']}).")
    elif brand and brand["primary"] and not colors & {brand["primary"], brand["accent"]}:
        notes.append(f"Note: the workspace has a brand kit (primary {brand['primary']}, accent "
                     f"{brand['accent']}) and this design uses neither — if it should be on-brand, fix that.")
    return spec, None, notes


def _prepare_pages(spec, brand, root):
    """A design of several pages, each prepared as a design of its own → (spec, error, notes)."""
    pages = spec.get("pages")
    if not isinstance(pages, list) or not pages:
        return None, ('Error: `pages` is a list of pages, each a design with a name: '
                      '[{"name": "Post", "size": "square", "fill": …, "nodes": […]}, …].'), []
    out, notes, seen = [], [], set()
    for n, page in enumerate(pages, 1):
        if not isinstance(page, dict):
            return None, f"Error: page {n} must be an object: {{name, size, fill, nodes}}.", []
        name = str(page.get("name") or "").strip()
        if not name:
            return None, f'Error: page {n} needs a `name` (e.g. "Post", "Story") — edits and downloads name a page.', []
        if name.lower() in seen:
            return None, f"Error: two pages are named {name!r} — every page needs its own name.", []
        seen.add(name.lower())
        if page.get("deck") is not None or page.get("pages") is not None:
            return None, (f"Error: page {name!r} is a design ({{size, fill, nodes}}) or a carousel "
                          f"({{frames}}) — a deck of layouts is a design of its own (render it by itself)."), []
        made, err, more = _prepare_spec({k: v for k, v in page.items() if k != "name"}, brand, root)
        if err:
            return None, err.replace("Error: ", f"Error: page {name!r}: ", 1), []
        out.append({"name": name[:40], **made})
        notes += [m for m in more if m not in notes]
    return {"pages": out}, None, notes


_DECK_THEMES = ("minimal-light", "minimal-dark", "bold-gradient", "editorial", "corporate", "tech-dark", "warm", "mono")
_DECK_KEYS = {"theme", "look", "size", "dir", "footer", "logo", "transition", "slides"}   # what layouts.js reads of a deck
# How a deck's slides or a document's pages are composed (cycls-design looks.js) — apart from the
# theme, its colours and faces. Without one, the look the theme is made for.
_LOOKS = ("classic", "editorial", "swiss", "bold", "soft", "minimal")
_THEME_LOOKS = {"editorial": "editorial", "tech-dark": "swiss", "warm": "soft", "mono": "minimal", "bold-gradient": "bold"}


def _look_for(theme):
    """The look a theme is made for: its own (a name) or its base's (overrides); else classic."""
    base = theme if isinstance(theme, str) else theme.get("base") if isinstance(theme, dict) else None
    return _THEME_LOOKS.get(str(base or "").strip().lower().replace(" ", "-").replace("_", "-"), "classic")


def _settle_look(spec, what, notes, default=True):
    """`spec`'s look checked and written plainly — or, with none and `default`, the theme's, written in
    (so a later change keeps it) and said. → an error, or None."""
    look = spec.get("look")
    if look is None:
        if default:
            spec["look"] = look = _look_for(spec.get("theme"))
            notes.append(f"The {what}'s look: {look}, its theme's — `look` composes it another way "
                         f"({', '.join(x for x in _LOOKS if x != look)}).")
        return None
    key = str(look).strip().lower() if isinstance(look, str) else ""
    if key not in _LOOKS:
        return f"Error: unknown look {json.dumps(look, ensure_ascii=False)} — a {what}'s look is one of: {', '.join(_LOOKS)}."
    spec["look"] = key
    return None


def _brand_theme(brand):
    """A minimal-light deck theme from a brand kit: primary hero, accent legible on white, its fonts."""
    primary, accent = brand.get("primary"), brand.get("accent") or brand.get("primary")
    theme = {"base": "minimal-light", "name": "brand"}
    hero_text = None
    if primary:
        hero_text = _readable_on(primary)
        theme.update(hero=primary, heroText=hero_text, heroMuted=hero_text + "b3")
    if accent:
        on_white = (accent if _contrast(accent, "#ffffff") >= 4.5
                    else primary if primary and _contrast(primary, "#ffffff") >= 4.5 else "#0f172a")
        theme.update(accent=on_white, accent2=accent, onAccent=_readable_on(on_white),
                     heroAccent=accent if not primary or _contrast(accent, primary) >= 3 else hero_text)
    if brand.get("heading"):
        theme["heading"] = f"{brand['heading']} Bold"
    if brand.get("body"):
        theme["body"] = brand["body"]
    return theme


def _prepare_deck(spec, brand, root):
    """A deck of layouts made ready for the service (cycls-design layouts.js): size resolved, image
    slots read, the brand kit as theme when there is one → (spec, error, notes)."""
    spec = json.loads(json.dumps(spec))
    deck = spec["deck"]
    slides = deck.get("slides")
    if not isinstance(slides, list) or not slides:
        return None, "Error: a deck needs `slides` — a list of {layout, …} (the layouts are in the tool description).", []
    try:
        deck["size"] = size = _frame_size({"size": deck.get("size") or "slide"})
    except ValueError as e:
        return None, f"Error: {e}.", []
    notes, total = [], 0
    moved = [k for k in sorted(_DECK_KEYS - {"slides"}) if spec.get(k) is not None and deck.get(k) is None]
    for k in moved:                                         # written beside the deck, not in it
        deck[k] = spec.pop(k)
    if moved:
        notes.append(f"{', '.join(f'`{k}`' for k in moved)} came beside `deck` and {'was' if len(moved) == 1 else 'were'} "
                     f"read as the deck's own — {'it goes' if len(moved) == 1 else 'they go'} inside it: "
                     f"{{\"deck\": {{\"theme\": …, \"slides\": […]}}}}.")
    unread = [k for k in deck if k not in _DECK_KEYS]
    if unread:
        notes.append(f"A deck doesn't read {', '.join(f'`{k}`' for k in unread)} — it reads "
                     f"{', '.join(sorted(_DECK_KEYS))}; colours and fonts go in `theme`.")
    theme = deck.get("theme")
    if isinstance(theme, dict):
        own = [k for k in theme if k not in ("base", "name", "logo")]
        notes.append(f"The deck's theme: {theme.get('base') or 'minimal-light'}"
                     f"{' with your ' + ', '.join(own) if own else ''}.")
    elif theme is None and not brand:
        notes.append("The deck has no `theme`, so it is minimal-light, the default — a theme of its own goes in the "
                     "deck: {\"deck\": {\"theme\": \"editorial\" or {\"base\": …, \"accent\": …, \"heading\": …}, \"slides\": […]}}.")

    def slot(value):
        nonlocal total
        resolved, n = _image_slot(value, root)
        total += n
        if total > _DESIGN_IMAGES_MAX:
            raise ValueError(f"the deck's images total over {_DESIGN_IMAGES_MAX >> 20} MB — use smaller copies "
                             f"(longest side ~2000px)")
        return resolved

    theme = deck.get("theme")
    if theme == "brand" and not brand:
        return None, (f"Error: theme \"brand\" needs a brand kit (brand/brand.yaml) — or pick a theme: "
                      f"{', '.join(_DECK_THEMES)}."), []
    if (theme is None or theme == "brand") and brand:
        deck["theme"] = theme = _brand_theme(brand)
        logo = next((f"brand/{n}" for n in _BRAND_LOGOS if root and (pathlib.Path(root) / "brand" / n).is_file()), None)
        if logo and not deck.get("logo"):
            deck["logo"] = logo
        fonts = " / ".join(f for f in (brand.get("heading"), brand.get("body")) if f)
        notes.append(f"The deck uses the workspace brand kit as its theme (primary {brand['primary']}, accent "
                     f"{brand['accent']}{', fonts ' + fonts if fonts else ''}{', logo ' + logo if logo else ''}).")
    if err := _settle_look(deck, "deck", notes):
        return None, err, []
    try:
        if deck.get("logo"):
            deck["logo"] = slot(deck["logo"])
        if isinstance(deck.get("footer"), dict) and deck["footer"].get("logo"):
            deck["footer"]["logo"] = slot(deck["footer"]["logo"])
        if isinstance(theme, dict) and theme.get("logo"):
            theme["logo"] = slot(theme["logo"])
        for i, s in enumerate(slides, 1):
            if not isinstance(s, dict):
                return None, f"Error: slide {i} must be an object: {{layout, …}}.", []
            if s.get("image"):
                s["image"] = slot(s["image"])
            for person in s.get("people") or []:
                if isinstance(person, dict) and person.get("photo"):
                    person["photo"] = slot(person["photo"])
            if isinstance(s.get("nodes"), list):
                sub, err, _ = _prepare_spec({"size": size, "fill": s.get("fill"), "nodes": s["nodes"]}, None, root)
                if err:
                    return None, err.replace("Error: ", f"Error: slide {i}: ", 1), []
                s["nodes"] = sub["nodes"]
    except ValueError as e:
        return None, f"Error: {e}.", []
    return spec, None, notes


_PAPER = {"a4": [1240, 1754], "letter": [1275, 1650], "a5": [874, 1240], "a3": [1754, 2480]}   # px at 150 dpi
# a document's own keys (document.js), read as its own when written beside it
_DOCUMENT_KEYS = ("title", "subtitle", "author", "date", "organization", "logo", "size", "landscape", "theme", "dir", "cover",
                  "columns", "openers", "toc", "numbered", "numbering", "figures", "header", "footer", "sections", "max_pages",
                  "kind", "meta", "contact", "back", "footnotes", "lang", "look")


_SHEET_ROWS = 2000      # max rows read from a spreadsheet


def _number(text):
    """A cell as a number — 1,204.5, 12%, (30) — or None."""
    if isinstance(text, bool):
        return None
    if isinstance(text, (int, float)):
        return text
    s = str(text or "").strip().replace(",", "").replace("\u00a0", "").replace(" ", "")
    if not s:
        return None
    negative = s.startswith("(") and s.endswith(")")
    s = s.strip("()").rstrip("%").lstrip("$€£")
    try:
        v = float(s)
    except ValueError:
        return None
    v = -v if negative else v
    return int(v) if v.is_integer() and "." not in s else v


def _sheet(src, root, sheet=None):
    """A workspace .csv / .tsv / .xlsx → (columns, rows, cell); ValueError with the fix."""
    if not isinstance(src, str) or not src.strip():
        raise ValueError('`from` is a workspace spreadsheet, e.g. "data/sales.csv"')
    path = _resolve_path(src, root)
    ext = path.suffix.lower()
    if ext not in (".csv", ".tsv", ".xlsx"):
        raise ValueError(f"{src} isn't a spreadsheet this reads — use a .csv, .tsv or .xlsx file")
    if not path.is_file():
        raise ValueError(f"{src} does not exist in the workspace")
    if ext == ".xlsx":
        try:
            import openpyxl
        except ImportError:
            raise ValueError(f"{src}: reading .xlsx needs openpyxl in this agent's image — save the sheet as .csv instead")
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            if sheet is not None and str(sheet) not in book.sheetnames:
                raise ValueError(f"{src} has no sheet {str(sheet)!r} — its sheets: {', '.join(book.sheetnames)}")
            ws = book[str(sheet)] if sheet is not None else book[book.sheetnames[0]]
            table = [list(r) for _, r in zip(range(_SHEET_ROWS + 1), ws.iter_rows(values_only=True))]
        finally:
            book.close()
    else:
        import csv
        text = path.read_text(encoding="utf-8-sig", errors="replace")
        table = [r for _, r in zip(range(_SHEET_ROWS + 1), csv.reader(text.splitlines(), delimiter="\t" if ext == ".tsv" else ","))]
    table = [r for r in table if any(c not in (None, "") for c in r)]
    if len(table) < 2:
        raise ValueError(f"{src} needs a header row and at least one row of data")
    cell = lambda v: "" if v is None else str(int(v)) if isinstance(v, float) and v.is_integer() else str(v).strip()
    columns = [cell(c) or f"Column {i + 1}" for i, c in enumerate(table[0])]
    rows = [[r[i] if i < len(r) else None for i in range(len(columns))] for r in table[1:]]
    return columns, rows, cell


def _block_data(b, root):
    """Fill a chart or table block that names a spreadsheet (`from`); ValueError with the fix."""
    chart, table = b.get("chart"), b.get("table")
    if isinstance(chart, dict) and chart.get("from"):
        columns, rows, cell = _sheet(chart["from"], root, chart.get("sheet"))

        def col(name):
            if str(name) not in columns:
                raise ValueError(f"{chart['from']} has no column {json.dumps(str(name))} — its columns: {', '.join(columns)}")
            return columns.index(str(name))
        x = col(chart["x"]) if chart.get("x") is not None else 0
        given = chart.get("y")
        ys = [col(y) for y in (given if isinstance(given, list) else [given])] if given is not None else [
            i for i in range(len(columns)) if i != x and rows and all(_number(r[i]) is not None for r in rows if r[i] not in (None, ""))
            and any(r[i] not in (None, "") for r in rows)]
        if not ys:
            raise ValueError(f"{chart['from']} has no column of numbers to chart — its columns: {', '.join(columns)}")
        series = []
        for i in ys:
            values = [_number(r[i]) for r in rows]
            if any(v is None for v in values):
                raise ValueError(f"{chart['from']}: column {json.dumps(columns[i])} isn't all numbers")
            series.append({"name": columns[i], "values": values})
        rest = {k: v for k, v in chart.items() if k not in ("from", "sheet", "x", "y")}
        b["chart"] = {**rest, "data": {"labels": [cell(r[x]) for r in rows], "series": series}}
    if isinstance(table, dict) and table.get("from"):
        columns, rows, cell = _sheet(table["from"], root, table.get("sheet"))
        pick = table.get("columns")
        if isinstance(pick, list) and pick:
            missing = [str(c) for c in pick if str(c) not in columns]
            if missing:
                raise ValueError(f"{table['from']} has no column {json.dumps(missing[0])} — its columns: {', '.join(columns)}")
            order = [columns.index(str(c)) for c in pick]
        else:
            order = list(range(len(columns)))
        limit = _num(table.get("limit"))
        if limit and limit > 0:
            rows = rows[:int(limit)]
        rest = {k: v for k, v in table.items() if k not in ("from", "sheet", "columns", "limit")}
        b["table"] = {**rest, "columns": [columns[i] for i in order], "rows": [[cell(r[i]) for i in order] for r in rows]}


def _prepare_document(spec, brand, root):
    """A document made ready for the service (cycls-design document.js / flow.js): images read,
    the brand kit as theme when there is one, hand-built nodes prepared → (spec, error, notes)."""
    spec = json.loads(json.dumps(spec))
    doc = spec["document"]
    moved = [k for k in _DOCUMENT_KEYS if spec.get(k) is not None and doc.get(k) is None]
    for k in moved:                                         # written beside the document, not in it
        doc[k] = spec.pop(k)
    sections = doc.get("sections")
    if not isinstance(sections, list) or not sections:
        return None, ("Error: a document needs `sections` — a list of {title, blocks: […]} "
                      "(the blocks are in the tool description)."), []
    if not str(doc.get("title") or "").strip():
        return None, "Error: a document needs a `title`.", []
    notes, total = [], 0
    if moved:
        notes.append(f"{', '.join(f'`{k}`' for k in moved)} came beside `document` and {'was' if len(moved) == 1 else 'were'} "
                     f"read as the document's own — {'it goes' if len(moved) == 1 else 'they go'} inside it: "
                     f"{{\"document\": {{\"title\": …, \"sections\": […]}}}}.")

    def slot(value):
        nonlocal total
        if isinstance(value, dict) and (isinstance(value.get("image"), str) or isinstance(value.get("svg"), str)):
            return value                                      # resolved already
        resolved, n = _image_slot(value, root)
        total += n
        if total > _DESIGN_IMAGES_MAX:
            raise ValueError(f"the document's images total over {_DESIGN_IMAGES_MAX >> 20} MB — use smaller copies "
                             f"(longest side ~2000px)")
        return resolved

    def nodes_of(nodes, size, fill=None):
        sub, err, _ = _prepare_spec({"size": size, "fill": fill, "nodes": nodes}, None, root)
        if err:
            raise ValueError(err.replace("Error: ", "", 1).rstrip("."))
        return sub["nodes"]

    def block(b):
        """One block's images and hand-built nodes, wherever they are nested."""
        if not isinstance(b, dict):
            return b
        for key in ("image", "photo", "picture"):
            if b.get(key):
                b[key] = slot(b[key])
        if str(b.get("type") or "") in ("image", "photo", "picture") and b.get("src") and not b.get("image"):
            b["image"] = slot(b.pop("src"))
        if isinstance(b.get("columns"), list):
            b["columns"] = [[block(x) for x in col] if isinstance(col, list) else block(col) for col in b["columns"]]
        if isinstance(b.get("nodes"), list):
            b["nodes"] = nodes_of(b["nodes"], [100000, 100000])
        _block_data(b, root)
        return b

    theme = doc.get("theme")
    if theme == "brand" and not brand:
        return None, (f"Error: theme \"brand\" needs a brand kit (brand/brand.yaml) — or pick a theme: "
                      f"{', '.join(_DECK_THEMES)}."), []
    if (theme is None or theme == "brand") and brand:
        doc["theme"] = _brand_theme(brand)
        logo = next((f"brand/{n}" for n in _BRAND_LOGOS if root and (pathlib.Path(root) / "brand" / n).is_file()), None)
        if logo and not doc.get("logo"):
            doc["logo"] = logo
        fonts = " / ".join(f for f in (brand.get("heading"), brand.get("body")) if f)
        notes.append(f"The document uses the workspace brand kit as its theme (primary {brand['primary']}, accent "
                     f"{brand['accent']}{', fonts ' + fonts if fonts else ''}{', logo ' + logo if logo else ''}).")
    # (Its look is settled before the source is kept — _render_document: a document rendered before
    # there were looks keeps the one it had through its section changes.)
    if err := _settle_look(doc, "document", notes, default=False):
        return None, err, []
    size = doc.get("size") or "a4"
    paper = list(_PAPER.get(str(size).lower().replace("-landscape", "").replace(" landscape", ""), _PAPER["a4"])) \
        if isinstance(size, str) else size
    try:
        if doc.get("logo"):
            doc["logo"] = slot(doc["logo"])
        if isinstance(doc.get("cover"), dict) and doc["cover"].get("image"):
            doc["cover"]["image"] = slot(doc["cover"]["image"])
        for i, s in enumerate(sections, 1):
            if not isinstance(s, dict):
                return None, f"Error: section {i} must be an object: {{title, blocks: […]}}.", []
            if isinstance(s.get("blocks"), list):
                s["blocks"] = [block(b) for b in s["blocks"]]
        for i, p in enumerate(doc.get("pages") or [], 1):
            if not isinstance(p, dict) or not isinstance(p.get("nodes"), list):
                return None, f"Error: pages[{i}] must be {{nodes: [spec nodes], fill?}} — a hand-built page.", []
            p["nodes"] = nodes_of(p["nodes"], paper, p.get("fill"))
    except ValueError as e:
        return None, f"Error: {e}.", []
    return spec, None, notes


_NO_FRAME = 100000              # the side of the stand-in frame a node added by an edit is prepared in
_PATTERN_OPACITY = 0.12         # a pattern is finish behind the text, unless it says otherwise

_OP_NUMERIC = ("x", "y", "dx", "dy", "w", "h", "size", "radius", "opacity", "letterSpacing",
               "lineHeight", "strokeWeight", "frame")


def _prepare_ops(ops, root):
    """Edit ops made safe to send: numbers read, replace_image files placed, added nodes and pages
    prepared as a render's are → (ops, error)."""
    if not isinstance(ops, list):
        return None, "Error: `ops` must be a list of operations."
    ops = json.loads(json.dumps(ops))
    for k, op in enumerate(ops, 1):
        if not isinstance(op, dict) or not op.get("op"):
            return None, f"Error: op {k} must be an object with an `op` (set_text, style, move, …)."
        for key in _OP_NUMERIC:
            if key in op and op[key] is not None:
                if (v := _num(op[key])) is None:
                    return None, f"Error: op {k}: `{key}` must be a number, not {op[key]!r}."
                op[key] = v
        if op["op"] == "replace_image":
            probe = {"src": op.pop("src", None), "w": 1, "h": 1}
            try:
                _place_image(probe, root)
            except ValueError as e:
                return None, f"Error: op {k}: {e}"
            op["image"] = probe["image"]
        elif op["op"] == "add":
            if not isinstance(op.get("node"), dict):
                return None, f"Error: op {k}: `add` needs a `node` (a spec node, e.g. a text or a stack)."
            spec, err, _ = _prepare_spec({"size": [_NO_FRAME, _NO_FRAME], "nodes": [op["node"]]}, None, root)
            if err:
                return None, err.replace("Error: ", f"Error: op {k}: ", 1)
            op["node"] = spec["nodes"][0]
        elif op["op"] == "page_add":
            if not isinstance(op.get("spec"), dict) or op["spec"].get("deck") or op["spec"].get("pages"):
                return None, (f"Error: op {k}: `page_add` needs `spec` — the page's design, "
                              f"{{size, fill, nodes}} (or a carousel's {{frames}}).")
            spec, err, _ = _prepare_spec(op["spec"], _load_brand(root) if root else None, root)
            if err:
                return None, err.replace("Error: ", f"Error: op {k}: ", 1)
            op["spec"] = spec
    return ops, None


def _mend_json_tail(text):
    """JSON valid up to its last value but closed with the wrong brackets → the object, re-closed;
    None if anything else is wrong (broken earlier, or cut off with no closing bracket)."""
    body = text.rstrip()
    tail = len(body)
    while tail and body[tail - 1] in "]} \t\r\n":
        tail -= 1
    if tail == len(body):
        return None                                           # no closing bracket: truncated
    stack, quoted, escaped = [], False, False
    for ch in body[:tail]:
        if quoted:
            if escaped:
                escaped = False
            elif ch == "\\":
                escaped = True
            elif ch == '"':
                quoted = False
        elif ch == '"':
            quoted = True
        elif ch in "{[":
            stack.append("}" if ch == "{" else "]")
        elif ch in "}]":
            if not stack or stack.pop() != ch:
                return None                                   # mismatched before the tail
    if quoted or not stack:
        return None
    try:
        return json.loads(body[:tail] + "".join(reversed(stack)))
    except json.JSONDecodeError:
        return None


def _prepare_slide(slide, settings, root):
    """One slide prepared as a deck's are → (slide, deck settings with images resolved, error)."""
    if not isinstance(slide, dict) or not (slide.get("layout") or isinstance(slide.get("nodes"), list)):
        return None, None, ('Error: `slide` is a layout slide ({"layout": "bullets", "title": …, "bullets": […]}) '
                            'or a hand-built one ({"nodes": […], "fill"?}).')
    spec, err, _ = _prepare_deck({"deck": {**settings, "slides": [slide]}}, None, root)
    if err:
        return None, None, err
    deck = spec["deck"]
    return deck.pop("slides")[0], deck, None
